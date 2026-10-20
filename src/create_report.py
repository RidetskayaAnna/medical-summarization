
from pathlib import Path

import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT


ROOT = Path(__file__).resolve().parent.parent
REPORTS = ROOT / "reports"
OUTPUT = REPORTS / "project_report.docx"


def read_csv(name):
    path = REPORTS / name
    if not path.exists():
        print(f"Предупреждение: не найден файл {path}")
        return pd.DataFrame()
    return pd.read_csv(path)


def set_cell(cell, value, bold=False):
    cell.text = str(value)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = bold
            run.font.size = Pt(9)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for cell, header in zip(table.rows[0].cells, headers):
        set_cell(cell, header, bold=True)

    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            set_cell(cell, value)

    doc.add_paragraph()
    return table


def add_bullet(doc, text):
    doc.add_paragraph(text, style="List Bullet")


def add_figure(doc, filename, caption):
    path = REPORTS / filename
    if path.exists():
        doc.add_picture(str(path), width=Inches(6.0))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        p = doc.add_paragraph(caption)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        doc.add_paragraph(f"График не найден: {filename}")


def fmt(value):
    try:
        return f"{float(value):.4f}"
    except (TypeError, ValueError):
        return str(value)


def main():
    REPORTS.mkdir(exist_ok=True)
    doc = Document()

    # Основное оформление
    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(12)
    styles["Normal"].paragraph_format.space_after = Pt(6)

    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        styles[style_name].font.name = "Times New Roman"
        styles[style_name].font.color.rgb = RGBColor(31, 78, 121)

    # Титульная страница
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("\n\n\nОТЧЁТ ПО ПРОЕКТУ\n").bold = True
    p.add_run(
        "Суммаризация историй болезни для передачи между врачами\n"
        "с использованием Zero-shot и Few-shot LLM"
    ).bold = True

    doc.add_paragraph("\n\n")
    p = doc.add_paragraph("Выполнила: __________________________")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("Группа: ______________________________")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("Год: 2026")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    # Аннотация
    doc.add_heading("Аннотация", level=1)
    doc.add_paragraph(
        "В проекте исследуется автоматическое формирование краткого "
        "структурированного резюме синтетической истории болезни с помощью "
        "большой языковой модели. Сравниваются извлекающий базовый метод, "
        "Zero-shot и Few-shot режимы модели FLAN-T5-base. Качество оценивается "
        "на отложенной тестовой выборке по ROUGE-L F1, приближенному Critical "
        "Fact Recall и автоматическому индикатору неподтверждённых чисел. "
        "Все записи являются искусственно сгенерированными учебными данными "
        "и не представляют реальные медицинские документы."
    )

    # 1. Постановка задачи
    doc.add_heading("1. Постановка задачи", level=1)
    doc.add_heading("1.1. Цель и проблема", level=2)
    doc.add_paragraph(
        "Цель проекта — сравнить способы автоматической суммаризации "
        "синтетических историй болезни и определить, насколько хорошо "
        "сгенерированное резюме сохраняет важные факты исходной записи. "
        "При передаче информации между специалистами пропуск симптома, "
        "результата измерения или документированного плана может изменить "
        "смысл записи. Поэтому одной лишь беглости текста недостаточно."
    )

    doc.add_heading("1.2. Формальное описание", level=2)
    add_bullet(doc, "Вход: текст синтетической истории болезни.")
    add_bullet(
        doc,
        "Выход: краткое резюме на естественном языке с симптомами, "
        "длительностью, результатами осмотра и измерений, а также "
        "документированным планом."
    )
    add_bullet(
        doc,
        "Ограничения: нельзя добавлять сведения, которых нет в источнике; "
        "числа, единицы измерения, результаты тестов и стороны тела должны "
        "сохраняться корректно."
    )
    add_bullet(
        doc,
        "В рамках учебного проекта система не предназначена для клинического "
        "использования и не заменяет проверку специалистом."
    )

    doc.add_heading("1.3. Требования", level=2)
    add_table(
        doc,
        ["Тип", "Требование"],
        [
            ("Функциональное", "Формировать краткое резюме исходного текста."),
            ("Функциональное", "Сохранять важные симптомы, измерения и план."),
            ("Функциональное", "Не добавлять неподтверждённые факты."),
            ("Функциональное", "Сравнивать базовый метод, Zero-shot и Few-shot."),
            ("Нефункциональное", "Повторяемость эксперимента и фиксированные разбиения."),
            ("Нефункциональное", "Воспроизводимая установка зависимостей и запуск."),
            ("Нефункциональное", "Использовать только синтетические данные."),
        ],
    )

    doc.add_heading("1.4. Метрики и критерий успеха", level=2)
    doc.add_paragraph(
        "Основная метрика — приближённый Critical Fact Recall: доля эталонных "
        "критических фактов, которые удалось обнаружить в сгенерированном "
        "резюме. Она выбрана потому, что для передачи медицинской информации "
        "важно не потерять значимые сведения. Реализация использует "
        "нормализацию и поиск текстовых совпадений, поэтому метрика является "
        "приближённой: перефразирование может не засчитаться, а семантическая "
        "ошибка не всегда будет обнаружена."
    )
    doc.add_paragraph(
        "Дополнительные метрики: ROUGE-L F1 для сравнения последовательностей "
        "токенов с эталонным резюме и доля примеров, где простой числовой "
        "проверяющий алгоритм обнаружил число, отсутствующее в источнике. "
        "Последний показатель является индикатором риска, а не полной проверкой "
        "фактической достоверности."
    )
    doc.add_paragraph(
        "В качестве предварительной исследовательской цели рассматривался "
        "Critical Fact Recall около 0,95. Это ориентир, а не подтверждённый "
        "порог клинической безопасности. Фактический результат оценивается "
        "по отложенной тестовой выборке."
    )

    # 2. Данные
    doc.add_heading("2. Датасет и подготовка данных", level=1)
    doc.add_heading("2.1. Паспорт датасета", level=2)
    doc.add_paragraph(
        "Датасет сгенерирован программно с фиксированным seed=42. Он содержит "
        "искусственные истории по пяти шаблонным сценариям: астма, "
        "гастроэнтерит, мигрень, пневмония и инфекция мочевыводящих путей. "
        "Записи не получены из реальных историй болезни."
    )
    add_table(
        doc,
        ["Поле", "Назначение"],
        [
            ("record_id", "Идентификатор синтетической записи."),
            ("scenario", "Сценарий, использованный при генерации."),
            ("source_text", "Исходная синтетическая история."),
            ("reference_summary", "Эталонное резюме."),
            ("critical_facts", "Список ключевых фактов в формате JSON-строки."),
            ("synthetic", "Признак синтетического происхождения записи."),
        ],
    )

    dataset_path = ROOT / "data" / "synthetic_medical_histories.csv"
    if dataset_path.exists():
        data = pd.read_csv(dataset_path)
        doc.add_paragraph(
            f"Исходный файл содержит {len(data)} записей. "
            f"Число столбцов: {len(data.columns)}."
        )
        if "scenario" in data.columns:
            counts = data["scenario"].value_counts()
            add_table(
                doc,
                ["Сценарий", "Количество"],
                [[str(k), int(v)] for k, v in counts.items()],
            )
        if "source_text" in data.columns:
            doc.add_paragraph(
                f"Уникальных исходных текстов: {data['source_text'].nunique()}. "
                f"Повторяющихся строк исходного текста удалено при подготовке "
                f"разбиений: {len(data) - data['source_text'].nunique()}."
            )

    doc.add_heading("2.2. Разбиение на выборки", level=2)
    add_table(
        doc,
        ["Выборка", "Количество", "Доля среди уникальных записей"],
        [
            ("Train", 130, "69,9%"),
            ("Validation", 28, "15,1%"),
            ("Test", 28, "15,1%"),
            ("Всего уникальных записей", 186, "100%"),
        ],
    )
    doc.add_paragraph(
        "Разбиение выполнено со стратификацией по сценариям и фиксированным "
        "seed=42 после исключения повторяющихся исходных текстов. Проверено "
        "отсутствие пересечения исходных текстов между выборками. Валидационная "
        "выборка использовалась для проверки вариантов, тестовая — для "
        "итогового сравнения зафиксированных методов."
    )

    doc.add_heading("2.3. Исследовательский анализ", level=2)
    eda_path = REPORTS / "eda_report.txt"
    if eda_path.exists():
        doc.add_paragraph(
            "Сводка из автоматически сформированного отчёта анализа данных:"
        )
        for line in eda_path.read_text(encoding="utf-8").splitlines():
            if line.strip():
                doc.add_paragraph(line.strip())
    else:
        doc.add_paragraph(
            "В проекте предусмотрен отдельный скрипт EDA. При анализе следует "
            "проверять распределение сценариев, пропущенные значения, дубликаты "
            "и размеры исходных текстов. Подробный вывод доступен в "
            "reports/eda_report.txt, если файл был создан."
        )

    add_figure(
        doc,
        "scenario_distribution.png",
        "Рисунок 1 — Распределение синтетических сценариев (если график создан).",
    )

    # 3. Методы
    doc.add_heading("3. Использованные методы", level=1)
    doc.add_heading("3.1. Базовый метод", level=2)
    doc.add_paragraph(
        "В качестве baseline использован простой извлекающий алгоритм: "
        "он выбирает предложения исходного текста по эвристикам, связанным "
        "с ключевыми словами и числовыми сведениями. Такой метод не использует "
        "обученную генеративную модель и позволяет получить отправную точку "
        "для сравнения."
    )

    doc.add_heading("3.2. Zero-shot", level=2)
    doc.add_paragraph(
        "Использована предварительно обученная модель google/flan-t5-base. "
        "В Zero-shot режиме модель получает инструкцию суммировать "
        "синтетический случай в 1–3 кратких предложениях, включив длительность "
        "и симптомы, данные осмотра, измерения и документированный план, "
        "не выдумывая информацию. Примеры правильных ответов в запрос не "
        "добавляются."
    )

    doc.add_heading("3.3. Few-shot", level=2)
    doc.add_paragraph(
        "В Few-shot режиме используется та же модель и основная инструкция, "
        "но в запрос добавляются два примера из обучающей выборки. Это позволяет "
        "проверить, улучшает ли демонстрация формата и примеров качество "
        "резюме. Для честного сравнения модель и тестовые записи остаются "
        "одинаковыми."
    )

    doc.add_heading("3.4. Параметры генерации", level=2)
    doc.add_paragraph(
        "Эксперимент выполнен в Google Colab на GPU Tesla T4 с моделью "
        "FLAN-T5-base в режиме FP16. Генерация выполнялась детерминированным "
        "beam search с num_beams=4, ограничением max_new_tokens=120, "
        "no_repeat_ngram_size=3 и repetition_penalty=1.2. Параметры были "
        "зафиксированы для тестового сравнения."
    )

    # 4. Результаты
    doc.add_heading("4. Результаты экспериментов", level=1)
    comparison = read_csv("method_comparison.csv")
    if not comparison.empty:
        headers = [
            "Метод",
            "ROUGE-L F1",
            "Critical Fact Recall",
            "Примеры с неподтверждёнными числами",
            "N",
        ]
        rows = []
        for _, r in comparison.iterrows():
            rows.append([
                r.get("method", ""),
                fmt(r.get("rougeL_f1", "")),
                fmt(r.get("critical_fact_recall", "")),
                (
                    "не рассчитывалось"
                    if pd.isna(r.get("unsupported_number_examples", float("nan")))
                    else f"{int(r['unsupported_number_examples'])}/{int(r.get('n_examples', 28))}"
                ),
                int(r.get("n_examples", 0)),
            ])
        add_table(doc, headers, rows)
    else:
        doc.add_paragraph(
            "Файл method_comparison.csv не найден. Создайте его скриптом "
            "сравнения методов перед повторной генерацией отчёта."
        )

    doc.add_paragraph(
        "На тестовой выборке Zero-shot показал наибольший Critical Fact Recall "
        "среди трёх методов — 0,7083, превысив baseline (0,6190). Few-shot "
        "получил наибольший ROUGE-L F1 (0,6620), однако его Critical Fact Recall "
        "составил 0,5952, то есть оказался ниже baseline и Zero-shot. "
        "Это показывает, что более близкое лексическое сходство с эталоном "
        "не обязательно означает лучшее сохранение критических фактов."
    )
    doc.add_paragraph(
        "Простой числовой проверяющий алгоритм обнаружил неподтверждённые "
        "числа в 9 из 28 Few-shot тестовых резюме и в 0 из 28 Zero-shot резюме. "
        "Нулевой результат Zero-shot не означает отсутствия всех фактических "
        "ошибок: алгоритм не обнаруживает замену одного числа другим, если "
        "новое число также встречается в исходном тексте, и не проверяет "
        "все семантические ошибки."
    )

    add_figure(
        doc,
        "method_comparison.png",
        "Рисунок 2 — Сравнение методов по ROUGE-L F1 и Critical Fact Recall.",
    )
    add_figure(
        doc,
        "unsupported_numbers.png",
        "Рисунок 3 — Количество тестовых примеров с обнаруженными неподтверждёнными числами.",
    )

    # 5. Анализ ошибок
    doc.add_heading("5. Анализ ошибок", level=1)
    doc.add_paragraph(
        "Для качественного анализа вручную рассмотрены 10 выбранных примеров "
        "из Few-shot тестовых результатов. Это целевая выборка характерных "
        "ошибок, а не случайная статистическая выборка и не независимая "
        "оценка несколькими экспертами."
    )

    reviewed = read_csv("error_analysis_reviewed.csv")
    if not reviewed.empty:
        preferred = [
            "record_id", "scenario", "error_type", "observation",
            "manual_label", "notes",
        ]
        available = [c for c in preferred if c in reviewed.columns]
        if available:
            rows = []
            for _, r in reviewed.iterrows():
                rows.append([
                    str(r.get(c, ""))[:500]
                    for c in available
                ])
            add_table(doc, available, rows)
        else:
            doc.add_paragraph(
                "Файл с ручным разбором найден, но ожидаемые колонки "
                "не обнаружены. См. reports/error_analysis_reviewed.csv."
            )
    else:
        doc.add_paragraph(
            "Файл error_analysis_reviewed.csv не найден. Подробный разбор "
            "можно посмотреть в reports/error_analysis.csv и в результатах "
            "эксперимента Few-shot."
        )

    doc.add_heading("5.1. Типичные категории ошибок", level=2)
    add_bullet(
        doc,
        "Искажение чисел: изменение длительности симптомов, количества "
        "эпизодов рвоты или жидкого стула, температуры либо сатурации."
    )
    add_bullet(
        doc,
        "Искажение клинической детали: изменение стороны расположения "
        "затемнения на снимке или описания признаков обезвоживания."
    )
    add_bullet(
        doc,
        "Пропуск важной информации: отсутствие документированного плана "
        "лечения или повторной оценки."
    )
    add_bullet(
        doc,
        "Неполное резюме: часть значимых измерений или результатов "
        "исходного текста не переносится в итог."
    )
    add_bullet(
        doc,
        "Неподходящий ответ: в отдельных случаях Zero-shot модель генерировала "
        "фразу вроде предложения написать историю вместо медицинского резюме."
    )
    doc.add_paragraph(
        "Ошибки показывают, что даже внешне связный текст может менять "
        "значимые факты. Поэтому ROUGE-L нельзя использовать как единственный "
        "критерий качества."
    )

    # 6. Ограничения
    doc.add_heading("6. Ограничения исследования", level=1)
    add_bullet(
        doc,
        "Использован небольшой синтетический датасет из 186 уникальных "
        "исходных текстов и пяти шаблонных сценариев; реальные медицинские "
        "документы не исследовались."
    )
    add_bullet(
        doc,
        "Тестовая выборка содержит 28 примеров, поэтому результаты могут "
        "существенно зависеть от состава выборки."
    )
    add_bullet(
        doc,
        "Синтетические записи генерируются по правилам и не отражают "
        "всей вариативности клинического языка."
    )
    add_bullet(
        doc,
        "Critical Fact Recall рассчитывается приблизительным текстовым "
        "сопоставлением, которое может не распознавать корректные перефразы."
    )
    add_bullet(
        doc,
        "Числовая проверка обнаруживает лишь часть потенциальных ошибок "
        "и не проверяет клиническую правдоподобность или все противоречия."
    )
    add_bullet(
        doc,
        "Few-shot использует только два примера; иные примеры и шаблоны "
        "могут изменить результаты."
    )
    add_bullet(
        doc,
        "Не проводилась клиническая валидация, оценка врачами в реальном "
        "рабочем процессе или проверка пригодности для принятия решений."
    )

    # 7. Выводы
    doc.add_heading("7. Выводы", level=1)
    doc.add_paragraph(
        "В проекте реализован воспроизводимый эксперимент по суммаризации "
        "синтетических историй болезни. Сравнены извлекающий baseline и два "
        "режима FLAN-T5-base — Zero-shot и Few-shot. На отложенной тестовой "
        "выборке Zero-shot показал лучший результат по основной метрике "
        "Critical Fact Recall (0,7083), тогда как Few-shot получил более "
        "высокий ROUGE-L F1 (0,6620), но меньший охват критических фактов "
        "(0,5952) и больше примеров с числами, не обнаруженными в источнике."
    )
    doc.add_paragraph(
        "Таким образом, добавление нескольких примеров в запрос не обеспечило "
        "улучшения сохранения критических фактов в данном эксперименте. "
        "Результаты подчёркивают необходимость проверки чисел, сторон тела, "
        "результатов исследований и документированного плана отдельно от "
        "лексического сходства. Работа является учебным исследованием на "
        "синтетических данных; её результаты не подтверждают клиническую "
        "безопасность или готовность к применению в медицине."
    )

    # 8. Воспроизводимость
    doc.add_heading("8. Воспроизводимость и структура проекта", level=1)
    doc.add_paragraph(
        "Код проекта организован в репозитории Git. Скрипты генерации данных, "
        "анализа, разбиения, baseline, сравнения, построения графиков и анализа "
        "ошибок находятся в каталоге src. Данные и результаты хранятся отдельно "
        "в data и reports. Эксперименты с FLAN-T5-base выполнялись в Google "
        "Colab, а итоговые CSV-файлы результатов переносились в локальный "
        "репозиторий."
    )
    doc.add_paragraph(
        "Для воспроизведения локальной части следует установить зависимости "
        "из requirements.txt и выполнять скрипты через Python из виртуального "
        "окружения. Для повторения эксперимента LLM необходим доступ к среде "
        "с установленными transformers, sentencepiece и PyTorch; конкретная "
        "доступность GPU зависит от среды запуска."
    )

    doc.add_heading("9. Источники и программные инструменты", level=1)
    add_bullet(doc, "Google FLAN-T5: google/flan-t5-base, модель для следования инструкциям.")
    add_bullet(doc, "ROUGE: метрика лексического сходства автоматических резюме с эталонами.")
    add_bullet(doc, "Python, pandas, scikit-learn, rouge-score, python-docx.")
    add_bullet(doc, "Git для контроля версий; Google Colab и PyTorch/Transformers для эксперимента LLM.")

    doc.save(OUTPUT)
    print(f"Отчёт создан: {OUTPUT}")
    print(f"Размер файла: {OUTPUT.stat().st_size} байт")


if __name__ == "__main__":
    main()
